"""원본/설정.xlsx 읽기 — 시스템에서 가져올 수 없어 사람이 정하는 값"""
from datetime import datetime

import openpyxl

USAGES = ("국내", "해외", "공용")


def _rows(ws):
    """머리글(3행) 아래의 비어 있지 않은 행"""
    for r in ws.iter_rows(min_row=4, values_only=True):
        r = [c.strip() if isinstance(c, str) else c for c in r]
        if any(c not in (None, "") for c in r):
            yield r


def _date(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str) and len(v) >= 10:
        return v[:10]
    return None


def load_settings(path):
    """(설정, 문제 목록)을 돌려준다. 문제는 대시보드에 그대로 표시한다."""
    wb = openpyxl.load_workbook(path, data_only=True)
    problems = []

    items = [{"code": r[0], "name": r[1], "cat": r[2]} for r in _rows(wb["관리품목"]) if r[0]]

    makers = {}
    for r in _rows(wb["제조사"]):
        code, _, maker, usage, qty, as_of = r[:6]
        if not code or not maker:
            continue
        if usage not in USAGES:
            problems.append(f"설정 '제조사' 시트: {code} {maker}의 용도 '{usage}'는 국내/해외/공용 중 하나여야 합니다")
            usage = "공용"
        makers.setdefault(code, []).append({
            "maker": maker, "usage": usage,
            "qty": qty if isinstance(qty, (int, float)) else None,
            "as_of": _date(as_of),
        })

    # 모델 코드에 특정 글자가 들어 있으면 그 제조사 재고를 먼저 쓴다 (예: TOPAZ → Crucial)
    model_maker = []
    if "모델제조사" in wb.sheetnames:
        for r in _rows(wb["모델제조사"]):
            if r[0] and r[1]:
                model_maker.append({"match": str(r[0]), "maker": r[1], "code": (r[2] if len(r) > 2 else None) or None})

    return {
        "items": items,
        "makers": makers,
        "model_maker": model_maker,
        "model_usage": {str(r[0]): r[1] for r in _rows(wb["모델용도"]) if r[0] and r[1] in ("국내", "해외")},
        # (거래처, 품목코드) → 제조사. 품목코드가 비어 있으면 그 거래처의 전체 품목에 적용
        "vendor_maker": {(r[0], (r[3] if len(r) > 3 else None) or None): r[1]
                         for r in _rows(wb["거래처"]) if r[0] and r[1]},
        "receipt_maker": {(str(r[0]), r[2]): r[6] for r in _rows(wb["입고제조사"])
                          if r[0] and len(r) > 6 and r[6]},
    }, problems
