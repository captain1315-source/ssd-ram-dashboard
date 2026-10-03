"""
원본 자료 → data/dashboard.json 변환 스크립트
사용법:
  python scripts/build_data.py

자료 출처
  시세            원본/Contract, Spot Price.xlsx (매일 루틴이 추가)
  재고·과부족     MRP 앱 DB (scm.py, 읽기 전용) + 원본/설정.xlsx (제조사·용도·기초수량)
  입고이력        원본/입고이력.csv (MRP 앱 DB의 새 입고를 실행할 때마다 누적, 원본/입력/의 ERP 파일도 흡수)
  견적 단가       원본/단가현황.xlsx
  시장 동향       원본/market.json (주간 루틴이 작성)
"""

import json
import sys
import warnings
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

import inventory as inv
import receipts as rc
import scm
from settings import load_settings
from update_prices import fix_and_load_xlsx

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "원본"
OUT = ROOT / "data" / "dashboard.json"
PRICE_XLSX = SRC / "Contract, Spot Price.xlsx"
SETTINGS_XLSX = SRC / "설정.xlsx"
QUOTES_XLSX = SRC / "단가현황.xlsx"
RECEIPTS_CSV = SRC / "입고이력.csv"
INBOX = SRC / "입력"
MARKET_JSON = SRC / "market.json"

SPOT_ITEMS = {
    "DDR4 8Gb (1Gx8) 3200": "ddr4_spot",
    "DDR5 16Gb (2Gx8) 4800/5600": "ddr5_spot",
    "512Gb TLC": "nand_spot",
}
CONTRACT_ITEMS = {
    "DDR4 8GB SO-DIMM": "ddr4_contract",
    "DDR5 16GB Module": "ddr5_contract",
}
TOL = 0.5          # 수량 비교 허용 오차
STALE_DAYS = 4     # MRP 계산이 이보다 오래되면 경고


def r2(v):
    if v is None:
        return None
    v = round(v, 2)
    return int(v) if v == int(v) else v


def comma(v):
    return f"{r2(v):,}"


def iso(d):
    return d.strftime("%Y-%m-%d")


def short_vendor(name):
    """'S18167_보이스아이주식회사' → '보이스아이주식회사'"""
    return name.split("_", 1)[1] if name and "_" in name else (name or "")


# ---------------------------------------------------------------- 시세

def read_prices(path):
    wb = fix_and_load_xlsx(str(path))
    out = {key: {} for key in list(SPOT_ITEMS.values()) + list(CONTRACT_ITEMS.values())}

    # Date | Category | Item | Daily High | Daily Low | Session High | Session Low | Session Avg | ...
    for r in wb["Spot Price"].iter_rows(min_row=2, values_only=True):
        if not isinstance(r[0], datetime) or r[2] not in SPOT_ITEMS or r[7] is None:
            continue
        # NAND Wafer는 Daily High/Low가 없어 Session High/Low를 쓴다
        high = r[3] if r[3] is not None else r[5]
        low = r[4] if r[4] is not None else r[6]
        out[SPOT_ITEMS[r[2]]][iso(r[0])] = {"d": iso(r[0]), "avg": r2(r[7]), "high": r2(high), "low": r2(low)}

    # Date | Mode | Item | High | Low | Average | High Chg | Low Chg | Avg Chg | History
    for r in wb["Contract Price"].iter_rows(min_row=2, values_only=True):
        if not isinstance(r[0], datetime) or r[2] not in CONTRACT_ITEMS or r[5] is None:
            continue
        out[CONTRACT_ITEMS[r[2]]][iso(r[0])] = {
            "d": iso(r[0]), "avg": r2(r[5]), "high": r2(r[3]), "low": r2(r[4]), "src": r[9],
        }

    # 같은 날짜가 여러 번 있으면 마지막 행을 쓴다
    return {key: [rows[d] for d in sorted(rows)] for key, rows in out.items()}


# ---------------------------------------------------------------- 입고이력 누적

def update_receipts(con, codes):
    """MRP 앱 DB와 입력 폴더의 새 입고를 누적 파일에 더하고 전체 목록을 돌려준다."""
    stored = rc.load(RECEIPTS_CSV)
    added = rc.merge(stored, rc.from_scm(scm.receipts(con, codes)))
    for path in sorted(INBOX.glob("*.xlsx")) if INBOX.exists() else []:
        if path.name.startswith("~$"):
            continue
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for ws in wb.worksheets:
            added += rc.merge(stored, rc.from_sheet(ws))
    if added:
        rc.save(RECEIPTS_CSV, stored)
    return stored, added


# ---------------------------------------------------------------- 재고·과부족

def by_maker(rows, cfg, code):
    """[(거래처, 수량, 입고번호, 비고)] → {제조사: 수량}. 제조사를 알 수 없으면 '미지정'."""
    out = defaultdict(float)
    for vendor, qty, no, note in rows:
        out[rc.resolve_maker(code, vendor, cfg, no, note) or inv.UNASSIGNED] += qty
    return dict(out)


def top_up(parts, total):
    """제조사별 합이 시스템 합계보다 적으면 차이를 '미지정'에 넣는다 (적송·기타입고 등 건별 자료가 없는 수량)."""
    gap = total - sum(parts.values())
    if gap > TOL:
        parts[inv.UNASSIGNED] = parts.get(inv.UNASSIGNED, 0) + gap
    return parts


def build_item(item, cfg, con, base, mrp_now, all_receipts, pos, netted, sop, issues):
    code = item["code"]
    m = mrp_now.get(code)
    if m is None:
        issues.append({"code": code, "msg": "MRP 앱에 이 품목의 계산 결과가 없습니다"})
        return None
    months = [scm.month_add(base, k) for k in range(7)]          # months[k] = M+k의 1일

    # ---- 품목 합계: MRP 앱 수치에서, 이미 들어온 것으로 확인된 발주(netted)만 입고 예정에서 뺀다
    cut = [sum(n["qty"] for n in netted if n["code"] == code and months[k] <= n["due"] < months[k + 1])
           for k in range(6)]
    backlog = max(m["po_backlog"] - cut[0], 0)
    demand = [m["month_demand"] + m["extra_demand"]] + [m[f"w{4 * k}_demand"] for k in range(1, 6)]
    inflow = [0] + [max(m[f"m{k}_in"] - cut[k], 0) for k in range(1, 6)]
    total = {
        "open": r2(m["carry_stock"]), "received": r2(m["receipt_qty"]), "pending": r2(backlog),
        "demand": [r2(v) for v in demand],
        "end": [r2(v) for v in inv.total_ends(m["carry_stock"], m["receipt_qty"], backlog, demand, inflow)],
    }

    # ---- 모델별 우선 제조사: 설정의 '모델제조사' 규칙 중 이 품목에 등록된 제조사만 적용
    makers = cfg["makers"].get(code, [])
    registered = {rc.norm_maker(mk["maker"]): mk["maker"] for mk in makers}
    rules = [(r["match"].lower(), registered[rc.norm_maker(r["maker"])]) for r in cfg["model_maker"]
             if r["code"] in (None, code) and rc.norm_maker(r["maker"]) in registered]
    prefer = lambda model: next((mk for text, mk in rules if text in (model or "").lower()), None)
    unknown_key = (inv.UNKNOWN, None)

    # ---- 소요 묶음(용도, 우선 제조사): 당월은 생산계획, M+1~M+5는 S&OP의 모델별 소요로 비율을 구해 합계에 적용
    raw, unknown = [], defaultdict(float)
    by_month = [scm.plan_demand_by_model(con, base, [code])[code]]
    by_month += [{model: ds[k] for model, ds in sop[code].items() if ds[k]} for k in range(5)]
    for models in by_month:
        s, u = inv.group_demand(models, cfg["model_usage"], prefer)
        raw.append(s)
        for token, qty in u.items():
            unknown[token] += qty
    fallback = defaultdict(float)
    for s in raw:
        for key, qty in s.items():
            fallback[key] += max(qty, 0)
    usage_demand = [inv.split_demand(demand[k], raw[k], fallback, unknown_key) for k in range(6)]
    if unknown:
        issues.append({"code": code, "msg":
                       "국내용인지 해외용인지 정해지지 않은 소요가 있습니다: "
                       + ", ".join(f"{t} {comma(q)}개" for t, q in sorted(unknown.items()))
                       + ". 설정의 '모델용도' 시트에 이 코드를 추가하고 국내/해외를 적어 주세요"})

    # ---- 제조사별 시작 재고: 설정의 기초수량에서 출발해 기준월 월초까지 이월
    as_of = max((mk["as_of"] for mk in makers if mk["as_of"]), default=None)
    if not makers or as_of is None or any(mk["qty"] is None for mk in makers):
        issues.append({"code": code, "msg": "설정의 '제조사' 시트에 기초수량과 기준일이 없어 제조사 구분 없이 표시합니다"})
        lots, as_of = [{"maker": inv.UNASSIGNED, "usage": "공용", "stock": m["carry_stock"]}], base
    else:
        lots = [{"maker": mk["maker"], "usage": mk["usage"], "stock": mk["qty"]} for mk in makers]
        if as_of[8:10] != "01":
            issues.append({"code": code, "msg": f"기초수량 기준일 {as_of}이 1일이 아니어서 그 달 월초 재고로 간주했습니다"})
        if as_of[:7] > base[:7]:
            issues.append({"code": code, "msg": f"기초수량 기준일 {as_of}이 MRP 기준월({base[:7]})보다 뒤입니다"})

    mine = [r for r in all_receipts if r["품번"] == code and r["입고일"] >= as_of]
    received_in = lambda month: by_maker(
        [(r["거래처"], r["수량"], r["입고번호"], r["비고"]) for r in mine if r["입고일"][:7] == month[:7]], cfg, code)

    month = as_of[:7] + "-01"
    while month < base:                                           # 기준일이 지난 달이면 그 사이를 이월
        past = scm.mrp(con, month, [code]).get(code)
        if past is None:
            issues.append({"code": code, "msg": f"{month[:7]} MRP 결과가 없어 그 달의 소진을 반영하지 못했습니다"})
        else:
            s, _ = inv.group_demand(scm.plan_demand_by_model(con, month, [code])[code], cfg["model_usage"], prefer)
            step = {"inflow": received_in(month),
                    "demand": inv.split_demand(past["month_demand"] + past["extra_demand"], s, fallback, unknown_key)}
            for lot, ends in zip(lots, inv.project(lots, [step])):
                lot["stock"] = ends[0]
        month = scm.month_add(month, 1)

    gap = sum(l["stock"] for l in lots) - m["carry_stock"]
    if abs(gap) > TOL:
        issues.append({"code": code, "msg": f"제조사별 재고 합계가 시스템 월초재고 {comma(m['carry_stock'])}와 "
                                            f"{comma(gap)} 차이 납니다. 설정의 기초수량을 새로 적어 주세요"})

    # ---- 제조사별 입고: 당월 입고(실적), 입고대기(당월 납기 발주잔량), M+1~M+5 발주
    received = top_up(received_in(base), m["receipt_qty"])
    my_pos = [p for p in pos if p["item_code"] == code]
    po_in = lambda k, qty: by_maker([(p["vendor_name"], qty(p), None, None) for p in my_pos
                                     if months[k] <= p["due"] < months[k + 1] and qty(p) > 0], cfg, code)
    pending = top_up(po_in(0, lambda p: p["order_qty"] - p["received_qty"]), backlog)
    future = [top_up(po_in(k, lambda p: p["order_qty"] - p["netted"]), inflow[k]) for k in range(1, 6)]

    known = {l["maker"] for l in lots}
    for name in sorted({mk for part in [received, pending, *future] for mk in part} - known):
        lots.append({"maker": name, "usage": "공용", "stock": 0})   # 미지정 입고는 공용으로 본다

    merged0 = defaultdict(float)
    for part in (received, pending):
        for name, qty in part.items():
            merged0[name] += qty
    steps = [{"inflow": dict(merged0) if k == 0 else future[k - 1], "demand": usage_demand[k]} for k in range(6)]
    ends = inv.project(lots, steps)

    out_lots = []
    for lot, end in zip(lots, ends):
        g, idx = inv.grade(end, None)
        out_lots.append({"maker": lot["maker"], "usage": lot["usage"], "open": r2(lot["stock"]),
                         "received": r2(received.get(lot["maker"], 0)), "pending": r2(pending.get(lot["maker"], 0)),
                         "end": [r2(v) for v in end], "grade": g, "short_idx": idx})

    # ---- 등급: 합계와 제조사 행 중 가장 나쁜 쪽. 용도가 다른 재고는 서로 대신 쓸 수 없다
    g_total, idx_total = inv.grade(total["end"], total["demand"])
    grade = inv.worst([g_total] + [l["grade"] for l in out_lots])
    short = next((l for l in out_lots if l["grade"] == grade and l["short_idx"] is not None), None)
    if g_total == grade and (idx_total is not None or short is None):
        short_idx, short_lot = idx_total, None
    else:
        short_idx, short_lot = short["short_idx"], {"maker": short["maker"], "usage": short["usage"]}

    return {
        "code": code, "name": item["name"], "cat": item["cat"],
        "grade": grade, "short_idx": short_idx, "short_lot": short_lot,
        "total": total, "lots": out_lots,
        # 용도별 월 소요. prefer가 있으면 그 제조사 재고를 먼저 쓰는 소요
        "usage_demand": [{"usage": u, "prefer": p, "demand": [r2(d.get((u, p), 0)) for d in usage_demand]}
                         for u, p in sorted({k for d in usage_demand for k in d},
                                            key=lambda k: (inv.USAGE_ORDER.index(k[0]), k[1] is not None, k[1] or ""))],
        "single_source": len([l for l in out_lots if l["maker"] != inv.UNASSIGNED]) == 1,
    }


def build_inventory(con, cfg, all_receipts, issues):
    codes = [it["code"] for it in cfg["items"]]
    base = scm.base_month(con)
    mrp_now = scm.mrp(con, base, codes)
    uploads = scm.upload_times(con)
    sop = scm.sop_by_model(con, codes)

    # 발주 자료는 입고 자료보다 드물게 올라온다. 발주 자료가 만들어진 날 이후의 입고는
    # 같은 품목·거래처의 미입고 발주에서 빼서, 이미 들어온 물량이 입고대기에 남지 않게 한다.
    po_at = uploads.get("발주_RAW", "")
    pos, netted = inv.net_received(scm.purchase_orders(con, codes), all_receipts, po_at[:10])
    items = [it for it in (build_item(it, cfg, con, base, mrp_now, all_receipts, pos, netted, sop, issues)
                           for it in cfg["items"]) if it]
    names = {it["code"]: it["name"] for it in cfg["items"]}
    adjustments = [f"{names[n['code']]}: {n['date']} 입고 {comma(n['qty'])}개를 {short_vendor(n['vendor'])} 발주"
                   f"(납기 {n['due']})의 입고대기에서 뺐습니다. 발주 자료({po_at})가 그 입고보다 먼저 올라와 있습니다"
                   for n in netted]

    computed = max((r["computed_at"] for r in mrp_now.values()), default="")[:16]
    if computed and (datetime.now() - datetime.fromisoformat(computed)).days > STALE_DAYS:
        issues.append({"code": None, "msg": f"MRP 앱의 마지막 계산이 {computed}입니다. 최신 자료가 올라갔는지 확인해 주세요"})

    months = [scm.month_add(base, k) for k in range(7)]
    open_pos = sorted(
        ({"code": p["item_code"], "due": p["due"], "vendor": short_vendor(p["vendor_name"]),
          "maker": rc.resolve_maker(p["item_code"], p["vendor_name"], cfg) or inv.UNASSIGNED,
          "qty": r2(p["order_qty"] - p["received_qty"]), "price": r2(p["unit_price"]), "currency": p["currency"]}
         for p in pos if p["order_qty"] - p["received_qty"] > 0 and months[0] <= p["due"] < months[6]),
        key=lambda p: (p["due"], p["code"]))

    return {
        "base_month": base[:7],
        "as_of": computed[:10], "computed_at": computed,
        "uploads": {k: uploads[k] for k in ("기초재고", "입고_RAW", "발주_RAW") if k in uploads},
        "months": [f"{int(months[k][5:7])}월말" for k in range(6)],
        "items": items,
        "pos": open_pos,
        "adjustments": adjustments,
        "issues": issues,
    }


# ---------------------------------------------------------------- 입고단가·견적

def read_quotes(path):
    """단가현황 목록(날짜·제조사·구분·스펙·단가) → 제조사 × 월 표. 같은 달에는 가장 늦은 날짜의 단가."""
    ws = openpyxl.load_workbook(path, data_only=True)["단가현황"]
    last, names = {}, {}
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if isinstance(r[0], datetime)]
    for d, maker, cat, spec, price, *_ in sorted(rows, key=lambda r: r[0]):
        if isinstance(price, (int, float)) and maker and spec:
            k = (cat, spec, rc.norm_maker(maker))
            names.setdefault(k, maker)
            last[(k, d.strftime("%Y-%m"))] = price
    months = sorted({m for _, m in last})
    return {"months": months,
            "rows": [{"cat": k[0], "spec": k[1], "maker": names[k], "p": [r2(last.get((k, m))) for m in months]}
                     for k in sorted(names)]}


def build_purchase(all_receipts, cfg, codes):
    slim = [(r["입고일"], r["품번"], r["통화"], r["수량"], r["금액"], r["단가"]) for r in all_receipts]
    usd = [r for r in all_receipts if r["품번"] in codes and r["통화"] == "USD" and r["수량"] and r["금액"]]
    recent = [{"d": r["입고일"], "code": r["품번"], "qty": r2(r["수량"]),
               "price": round(rc.unit_price(r["수량"], r["금액"], r["단가"]), 2),
               "maker": rc.resolve_maker(r["품번"], r["거래처"], cfg, r["입고번호"], r["비고"]),
               # 기초수량 기준일 이후 입고만 제조사 지정이 필요하다 (그 전 입고는 기초수량에 이미 포함)
               "needs_maker": r["입고일"] >= baseline(cfg, r["품번"]),
               "vendor": short_vendor(r["거래처"]), "no": r["입고번호"]}
              for r in sorted(usd, key=lambda r: r["입고일"])[-10:]][::-1]
    issues = [{"code": code, "msg": f"입고이력 {d}: 수량 {comma(qty)} × 단가 {comma(price)} ≠ 금액 {comma(amt)}"}
              for d, code, qty, price, amt in rc.amount_mismatches(slim, codes)]
    return {"receipts": rc.monthly_receipts(slim, codes), "recent": recent,
            "quotes": read_quotes(QUOTES_XLSX), "issues": issues}


def baseline(cfg, code):
    """품목의 기초수량 기준일. 없으면 어떤 날짜보다도 뒤로 취급되는 값."""
    return max((mk["as_of"] for mk in cfg["makers"].get(code, []) if mk["as_of"]), default="9999")


def unassigned_receipts(all_receipts, cfg):
    """기초수량 기준일 이후 입고 중 제조사를 정할 수 없는 건"""
    out = []
    for r in all_receipts:
        if r["입고일"] >= baseline(cfg, r["품번"]) and \
                not rc.resolve_maker(r["품번"], r["거래처"], cfg, r["입고번호"], r["비고"]):
            out.append({"code": r["품번"], "msg":
                        f"제조사 미지정 입고: {r['입고일']} {short_vendor(r['거래처'])} {comma(r['수량'])}개 "
                        f"(입고번호 {r['입고번호']}). 설정의 '거래처' 시트에 품목별 제조사를 적거나 '입고제조사' 시트에 건별로 적어 주세요"})
    return out


# ---------------------------------------------------------------- main

def main():
    for path in (PRICE_XLSX, SETTINGS_XLSX, QUOTES_XLSX):
        if not path.exists():
            sys.exit(f"파일 없음: {path}")

    cfg, problems = load_settings(SETTINGS_XLSX)
    codes = [it["code"] for it in cfg["items"]]
    issues = [{"code": None, "msg": p} for p in problems]

    con = scm.connect()
    all_receipts, added = update_receipts(con, codes)
    issues += unassigned_receipts(all_receipts, cfg)
    inventory = build_inventory(con, cfg, all_receipts, issues)
    con.close()

    prices = read_prices(PRICE_XLSX)
    market = json.loads(MARKET_JSON.read_text(encoding="utf-8")) if MARKET_JSON.exists() else None
    data = {"prices": prices, "inventory": inventory,
            "purchase": build_purchase(all_receipts, cfg, codes), "market": market}
    text = json.dumps(data, ensure_ascii=False, indent=1)

    OUT.parent.mkdir(exist_ok=True)
    changed = not OUT.exists() or OUT.read_text(encoding="utf-8") != text
    if changed:
        OUT.write_text(text, encoding="utf-8")

    print(f"시세 기준일: { {k: v[-1]['d'] for k, v in prices.items() if v} }")
    print(f"재고: MRP 기준월 {inventory['base_month']}, 계산 {inventory['computed_at']}")
    for it in inventory["items"]:
        print(f"  {it['code']} {it['grade']:8s} 합계 월말 {it['total']['end']}")
        for l in it["lots"]:
            print(f"      {l['maker']:14s} {l['usage']} 월초 {l['open']} 입고 {l['received']} 대기 {l['pending']} → {l['end']}")
    print(f"입고이력: {len(all_receipts)}건 (이번에 {added}건 추가)")
    for i in inventory["issues"] + data["purchase"]["issues"]:
        print(f"  [확인] {i['code'] or '-'} {i['msg']}")
    print(f"시장 동향: {market['updated'] if market else '없음'}")
    print(f"{'저장' if changed else '변경 없음'}: {OUT}")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
